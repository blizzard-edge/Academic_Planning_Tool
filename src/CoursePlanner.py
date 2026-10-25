
"""SCPTAA_001: Semester-by-semester academic course planner.

Produces exporter-compatible semester plans using catalog prerequisites,
DegreeWorks progress, and confirmed course offerings.

Unverified grades and expected completion of in-progress prerequisites
are explicitly identified as conditional recommendations.
"""

import re

from openpyxl import load_workbook

# The dedicated checker owns prerequisite interpretation and evaluation.
import PrerequisiteChecker as prerequisite_checker
from PrerequisiteChecker import (
    UnsupportedPrerequisiteError,
    normalize_course_code,
    normalize_course_list,
    tokenize_prerequisites,
    PrerequisiteExpressionParser,
    parse_prerequisite_details,
    parse_prerequisite_groups,
    prerequisites_satisfied,
    _student_grades,
    _evaluate_requirements,
)


SEMESTER_ORDER = {"spring": 0, "summer": 1, "fall": 2}
TERM_PREFIXES = {"SP": "Spring", "SU": "Summer", "FA": "Fall"}
CONFIRMED_DELIVERY_CODES = {"F", "O", "D", "N"}
COURSE_CODE_PATTERN = re.compile(r"[A-Z]{4}\s+\d{4}[A-Z]?")
def extract_credits(course_hours):
    if isinstance(course_hours, int) and not isinstance(course_hours, bool):
        return course_hours if course_hours >= 0 else None
    if not isinstance(course_hours, str):
        return None
    match = re.fullmatch(
        r"\(\s*\d+\s*-\s*\d+\s*-\s*(\d+)\s*\)",
        course_hours.strip(),
    )
    return int(match.group(1)) if match else None


def semester_sort_key(semester):
    if not isinstance(semester, str):
        raise ValueError("Semester must be a string.")
    match = re.fullmatch(
        r"\s*(Spring|Summer|Fall)\s+(\d{4})\s*",
        semester,
        re.IGNORECASE,
    )
    if not match:
        raise ValueError(
            f"Invalid semester '{semester}'. Expected 'Spring 2027'."
        )
    return int(match.group(2)), SEMESTER_ORDER[match.group(1).lower()]


def generate_semesters(start_semester, count):
    if not isinstance(count, int) or isinstance(count, bool) or count < 1:
        raise ValueError("Semester count must be positive.")
    year, season = semester_sort_key(start_semester)
    seasons = ["Spring", "Summer", "Fall"]
    return [
        f"{seasons[(season + i) % 3]} {year + (season + i) // 3}"
        for i in range(count)
    ]


def parse_schedule_heading(heading):
    if not isinstance(heading, str):
        return None
    match = re.fullmatch(r"(SP|SU|FA)(\d{2})", heading.strip().upper())
    if not match:
        return None
    return f"{TERM_PREFIXES[match.group(1)]} {2000 + int(match.group(2))}"


def is_discontinued_cell(cell):
    fill = cell.fill
    return (
        fill.patternType == "solid"
        and fill.fgColor.type == "theme"
        and fill.fgColor.theme == 1
        and abs(fill.fgColor.tint - 0.5) < 0.01
    )


def has_confirmed_offering(value):
    if not isinstance(value, str):
        return False
    value = value.strip().upper()
    if not value or "?" in value:
        return False
    codes = [part.strip() for part in value.split(",")]
    return all(code in CONFIRMED_DELIVERY_CODES for code in codes)


def load_course_offerings(workbook_path, start_semester="Fall 2026"):
    """Read only confirmed future offerings from the schedule workbook."""
    start_key = semester_sort_key(start_semester)
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        sheet = workbook.active
        columns = {}
        for column in range(3, sheet.max_column + 1):
            term = parse_schedule_heading(sheet.cell(3, column).value)
            if term and semester_sort_key(term) >= start_key:
                columns[column] = term

        offerings = {}
        for row in range(4, sheet.max_row + 1):
            code = normalize_course_code(sheet.cell(row, 1).value)
            if not COURSE_CODE_PATTERN.fullmatch(code):
                continue
            for column, term in columns.items():
                cell = sheet.cell(row, column)
                if is_discontinued_cell(cell):
                    continue
                if has_confirmed_offering(cell.value):
                    terms = offerings.setdefault(code, [])
                    if term not in terms:
                        terms.append(term)
        return offerings
    finally:
        workbook.close()


def build_course_index(course_catalog):
    return {
        normalize_course_code(record["course_code"]): record
        for record in course_catalog
        if isinstance(record, dict)
        and normalize_course_code(record.get("course_code"))
    }


def get_course_credits(course, default_credits=None):
    credits = course.get("credits")
    if isinstance(credits, int) and not isinstance(credits, bool) and credits >= 0:
        return credits
    parsed = extract_credits(course.get("course_hours"))
    return parsed if parsed is not None else default_credits


def create_course_entry(code, catalog_record, credits):
    return {
        "course_code": code,
        "title": catalog_record.get("title", ""),
        "credits": credits,
        "notes": "",
    }


def build_course_plan(
    required_courses,
    degreeworks_data,
    course_catalog,
    offering_schedule,
    start_semester="Fall 2026",
    max_credits=12,
    include_summer=True,
    default_credits=None,
    current_semester="Fall 2026",
):
    """Create an exporter-compatible course plan and report limitations."""
    if not isinstance(max_credits, int) or isinstance(max_credits, bool) or max_credits <= 0:
        raise ValueError("max_credits must be a positive integer.")
    if not isinstance(include_summer, bool):
        raise ValueError("include_summer must be a boolean.")
    if not isinstance(required_courses, list):
        raise ValueError("required_courses must be a list.")
    if not isinstance(degreeworks_data, dict):
        raise ValueError("degreeworks_data must be a dictionary.")
    required_fields = ("completed_courses", "in_progress_courses")
    missing = [field for field in required_fields if field not in degreeworks_data]
    if missing:
        raise ValueError("Invalid DegreeWorks data: missing " + ", ".join(missing))
    for field in required_fields:
        if not isinstance(degreeworks_data[field], list):
            raise ValueError(f"Invalid DegreeWorks data: {field} must be a list.")
    if not isinstance(course_catalog, list):
        raise ValueError("course_catalog must be a list.")
    if not isinstance(offering_schedule, dict):
        raise ValueError("offering_schedule must be a dictionary.")

    start_key = semester_sort_key(start_semester)
    current_key = semester_sort_key(current_semester)

    completed_records = degreeworks_data.get("completed_courses", [])
    in_progress_records = degreeworks_data.get("in_progress_courses", [])

    completed = normalize_course_list(completed_records)
    in_progress = normalize_course_list(in_progress_records) - completed
    grades = _student_grades(completed_records)
    catalog_index = build_course_index(course_catalog)

    offerings = {}
    all_terms = set()

    for code, terms in offering_schedule.items():
        code = normalize_course_code(code)
        if not code or not isinstance(terms, list):
            continue
        for term in terms:
            try:
                key = semester_sort_key(term)
            except ValueError:
                continue
            if key < start_key:
                continue
            if not include_summer and key[1] == SEMESTER_ORDER["summer"]:
                continue
            offerings.setdefault(code, set()).add(term)
            all_terms.add(term)

    semesters = sorted(all_terms, key=semester_sort_key)
    plan = [{"semester": term, "courses": []} for term in semesters]

    pending = []
    seen = set()
    for item in required_courses:
        code = normalize_course_code(
            item.get("course_code") if isinstance(item, dict) else item
        )
        if not code or code in seen:
            continue
        seen.add(code)
        if code not in completed and code not in in_progress:
            pending.append(code)

    parsed = {}
    errors = {}
    for code in pending:
        record = catalog_index.get(code)
        if record is None:
            continue
        try:
            parsed[code] = prerequisite_checker.parse_prerequisite_details(record.get("prerequisites"))
        except UnsupportedPrerequisiteError as error:
            errors[code] = str(error)

    scheduled = set()
    unscheduled = []
    warnings = []
    conditional_count = 0

    for entry in plan:
        term = entry["semester"]
        term_key = semester_sort_key(term)
        credits_used = 0

        # A course taken during this semester cannot satisfy an earlier
        # prerequisite unless explicit concurrency support is implemented.
        available = completed | scheduled
        allow_expected = term_key > current_key
        if allow_expected:
            available |= in_progress

        for code in list(pending):
            if term not in offerings.get(code, set()):
                continue
            record = catalog_index.get(code)
            if record is None or code in errors:
                continue

            credits = get_course_credits(record, default_credits)
            if credits is None or credits_used + credits > max_credits:
                continue

            expression, minimum_grades, concurrent = parsed[code]

            eligible, notes = prerequisite_checker._evaluate_requirements(
                expression,
                available,
                completed,
                in_progress,
                grades,
                minimum_grades,
                allow_expected,
            )
            if not eligible:
                continue

            # Concurrency is recognized, but the scheduler does not yet
            # approve same-term enrollment automatically.
            if concurrent:
                notes.append(
                    "Concurrent enrollment policy requires advisor verification"
                )

            course_entry = create_course_entry(code, record, credits)
            if notes:
                course_entry["notes"] = "CONDITIONAL: " + "; ".join(
                    dict.fromkeys(notes)
                )
                conditional_count += 1

            entry["courses"].append(course_entry)
            credits_used += credits
            scheduled.add(code)
            pending.remove(code)

    for code in pending:
        record = catalog_index.get(code)
        if record is None:
            reason = "Course not found in the course catalog."
        elif code in errors:
            reason = "Prerequisite requires advisor review: " + errors[code]
        elif get_course_credits(record, default_credits) is None:
            reason = "Credit hours could not be determined."
        elif not offerings.get(code):
            reason = "No confirmed course offerings found."
        else:
            expression, minimum_grades, _ = parsed[code]
            possible, _ = prerequisite_checker._evaluate_requirements(
                expression,
                completed | scheduled | in_progress,
                completed,
                in_progress,
                grades,
                minimum_grades,
                True,
            )
            if not possible:
                reason = "Prerequisites are not satisfied."
            else:
                reason = (
                    "Could not fit course into an available semester "
                    "within the credit limit and planning window."
                )

        unscheduled.append({"course_code": code, "reason": reason})

    if in_progress:
        warnings.append(
            "In-progress courses are excluded from rescheduling. "
            "Future prerequisite eligibility assumes successful completion "
            "and must be verified."
        )
    if conditional_count:
        warnings.append(
            f"{conditional_count} scheduled course(s) have conditional "
            "prerequisite requirements requiring verification."
        )
    if unscheduled:
        warnings.append(
            f"{len(unscheduled)} required course(s) could not be scheduled."
        )

    return {
        "plan": plan,
        "unscheduled": unscheduled,
        "warnings": warnings,
    }
