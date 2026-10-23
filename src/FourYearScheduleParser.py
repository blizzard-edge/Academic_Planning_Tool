import os
import re
import sys

from openpyxl import load_workbook


COURSE_CODE_PATTERN = re.compile(r"^[A-Z]{4} \d{4}[A-Z]?$")

TERM_PATTERN = re.compile(r"^(SP|SU|FA)(\d{2})$")
SEASON_NAMES = {"SP": "Spring", "SU": "Summer", "FA": "Fall"}

OFFERED_PATTERN = re.compile(r"^\(?\s*[FODN](?![A-Za-z])")


class ScheduleFormatError(ValueError):
    """Raised when the four-year schedule file is not in the expected layout."""


def _clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _clean_course_code(value):
    text = _clean_text(value).upper()
    if COURSE_CODE_PATTERN.match(text):
        return text
    return None


def _is_greyed(cell):
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False

    color = fill.fgColor
    if color is None:
        return False

    if color.type == "theme":
        tint = color.tint or 0
        return color.theme in (0, 1) and tint != 0

    if color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) == 8:
        red, green, blue = color.rgb[2:4], color.rgb[4:6], color.rgb[6:8]
        return red == green == blue and red not in ("FF", "00")

    return False


def _open_workbook(path):
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    if not path.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Four-year schedule input must be an Excel (.xlsx) file.")

    try:
        return load_workbook(path, data_only=True)
    except Exception as error:
        raise ValueError(f"Could not read Excel file '{path}': {error}") from error


def _find_header_row(sheet):
    for row in range(1, min(sheet.max_row, 20) + 1):
        if _clean_text(sheet.cell(row, 1).value).lower() == "course":
            return row
    return None


def load_four_year_schedule(path):
    workbook = _open_workbook(path)

    sheet = None
    header_row = None
    for candidate in workbook.worksheets:
        header_row = _find_header_row(candidate)
        if header_row:
            sheet = candidate
            break

    if sheet is None:
        raise ScheduleFormatError(
            "Could not find the schedule header row (a 'Course' cell in column A)."
        )

    terms = []
    term_columns = {}
    title_column = 2

    for cell in sheet[header_row]:
        text = _clean_text(cell.value)

        if text.lower() == "course title":
            title_column = cell.column
            continue

        match = TERM_PATTERN.match(text.upper())
        if match:
            code = text.upper()
            terms.append({
                "code": code,
                "season": SEASON_NAMES[match.group(1)],
                "year": 2000 + int(match.group(2)),
            })
            term_columns[code] = cell.column

    if not term_columns:
        raise ScheduleFormatError(
            "Could not find any term columns (SP/SU/FA + year) in the header row."
        )

    last_term_column = max(term_columns.values())

    courses = {}

    for row in range(header_row + 1, sheet.max_row + 1):
        course_code = _clean_course_code(sheet.cell(row, 1).value)

        if course_code is None:
            continue

        if course_code in courses:
            raise ScheduleFormatError(
                f"Course {course_code} appears more than once in the schedule (row {row})."
            )

        term_data = {}
        for term_code, column in term_columns.items():
            cell = sheet.cell(row, column)
            raw = _clean_text(cell.value) or None
            greyed = _is_greyed(cell)

            term_data[term_code] = {
                "raw": raw,
                "offered": (
                    raw is not None
                    and not greyed
                    and OFFERED_PATTERN.match(raw) is not None
                ),
                "greyed": greyed,
            }

        notes = []
        for column in range(last_term_column + 1, sheet.max_column + 1):
            text = _clean_text(sheet.cell(row, column).value)
            if text:
                notes.append(text)

        courses[course_code] = {
            "course_code": course_code,
            "title": _clean_text(sheet.cell(row, title_column).value),
            "notes": notes,
            "terms": term_data,
        }

    if not courses:
        raise ScheduleFormatError("No course rows were found in the schedule.")

    return {"terms": terms, "courses": courses}


def get_offered_terms(schedule, course_code):
    code = _clean_course_code(course_code)

    if code is None or code not in schedule["courses"]:
        raise KeyError(f"Course not found in the four-year schedule: {course_code}")

    return [
        term["code"]
        for term in schedule["terms"]
        if schedule["courses"][code]["terms"][term["code"]]["offered"]
    ]


def is_offered(schedule, course_code, term_code):
    code = _clean_course_code(course_code)

    if code is None or code not in schedule["courses"]:
        raise KeyError(f"Course not found in the four-year schedule: {course_code}")

    term = schedule["courses"][code]["terms"].get(str(term_code).strip().upper())

    if term is None:
        raise KeyError(f"Term not found in the four-year schedule: {term_code}")

    return term["offered"]


def print_schedule_results(schedule, course_codes=None):
    print("\nFOUR-YEAR SCHEDULE: TERMS OFFERED")
    print("---------------------------------")

    if course_codes is None:
        course_codes = list(schedule["courses"])

    for course_code in course_codes:
        try:
            offered = get_offered_terms(schedule, course_code)
        except KeyError as error:
            print(error.args[0])
            continue

        title = schedule["courses"][_clean_course_code(course_code)]["title"]
        print(course_code, "|", title, "|", ", ".join(offered) if offered else "None")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python src/FourYearScheduleParser.py <four_year_schedule_xlsx> [course_code ...]")
        print('Example: python src/FourYearScheduleParser.py schedule.xlsx "CPSC 6109"')
        sys.exit(1)

    schedule_path = sys.argv[1]
    requested_courses = sys.argv[2:] or None

    try:
        schedule_data = load_four_year_schedule(schedule_path)
        print_schedule_results(schedule_data, requested_courses)

    except Exception as error:
        print(f"Error reading four-year schedule: {error}")
        sys.exit(1)
