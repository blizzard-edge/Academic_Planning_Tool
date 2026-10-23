import os
import re
import sys

from openpyxl import load_workbook



COURSE_CODE_PATTERN = re.compile(r"^[A-Z]{4} \d{4}[A-Z]?$")

SEASON_WORDS = ("fall", "spring", "summer")


class StudyPlanFormatError(ValueError):
    """Raised when the study plan file is not in the expected layout."""


def _clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _clean_course_code(value):
    text = _clean_text(value).upper()
    if COURSE_CODE_PATTERN.match(text):
        return text
    return None


def _open_workbook(path):
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    if not path.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Graduate Study Plan input must be an Excel (.xlsx) file.")

    try:
        return load_workbook(path, data_only=True)
    except Exception as error:
        raise ValueError(f"Could not read Excel file '{path}': {error}") from error



def _find_course_info_sheet(workbook):
    for sheet in workbook.worksheets:
        for row in range(1, min(sheet.max_row, 10) + 1):
            for cell in sheet[row]:
                if _clean_text(cell.value).lower() == "course name":
                    return sheet, row
    return None, None


def parse_course_information(workbook):
    sheet, header_row = _find_course_info_sheet(workbook)

    if sheet is None:
        raise StudyPlanFormatError(
            "Could not find the course information sheet (header 'Course name')."
        )

    columns = {}
    for cell in sheet[header_row]:
        text = _clean_text(cell.value).lower()
        if text:
            columns[text] = cell.column

    for required in ("course name", "course number"):
        if required not in columns:
            raise StudyPlanFormatError(
                f"Course information sheet is missing the '{required}' column."
            )

    def read(row, header):
        column = columns.get(header)
        return _clean_text(sheet.cell(row, column).value) if column else ""

    courses = {}

    for row in range(header_row + 1, sheet.max_row + 1):
        course_code = _clean_course_code(sheet.cell(row, columns["course number"]).value)

        if course_code is None:
            continue

        required_in = [
            text for text in (read(row, "required in #1"), read(row, "required in #2"))
            if text
        ]
        offered = [part.strip() for part in read(row, "offer").split(",") if part.strip()]

        courses[course_code] = {
            "course_code": course_code,
            "course_name": read(row, "course name"),
            "prerequisite": read(row, "prerequisite") or None,
            "required_in": required_in,
            "offered": offered,
            "instructor": read(row, "instructor") or None,
        }

    return courses



def _find_plan_sheet(workbook):
    for sheet in workbook.worksheets:
        for row in range(1, min(sheet.max_row, 30) + 1):
            header_columns = {}
            for cell in sheet[row]:
                text = _clean_text(cell.value)
                if text.lower().startswith(SEASON_WORDS):
                    header_columns[cell.column] = text
            if len(header_columns) >= 3:
                return sheet, row, header_columns
    return None, None, None


def parse_study_plans(workbook):
    sheet, header_row, header_columns = _find_plan_sheet(workbook)

    if sheet is None:
        raise StudyPlanFormatError(
            "Could not find the study plan grid (a header row with Fall / Spring / Summer)."
        )

    semester_columns = sorted(header_columns)
    label_column = semester_columns[0] - 1

    if label_column < 1:
        raise StudyPlanFormatError("Study plan grid has no label column before the semester columns.")

    last_semester_column = semester_columns[-1]

    plans = []
    unassigned_cells = []
    notes = []

    current_track = None
    current_plan = None
    current_start_column = None

    for row in range(header_row + 1, sheet.max_row + 1):
        label = _clean_text(sheet.cell(row, label_column).value)

        cells = {}
        for column in semester_columns:
            text = _clean_text(sheet.cell(row, column).value)
            if text:
                cells[column] = text

        for column in range(last_semester_column + 1, sheet.max_column + 1):
            text = _clean_text(sheet.cell(row, column).value)
            if text:
                notes.append({"cell": sheet.cell(row, column).coordinate, "text": text})

        if label:
            lowered = label.lower()

            if lowered.endswith("start") and lowered.startswith(SEASON_WORDS):
                start_season = lowered.split()[0]
                start_column = next(
                    (c for c in semester_columns
                     if header_columns[c].lower().startswith(start_season)),
                    None,
                )

                if start_column is None:
                    raise StudyPlanFormatError(
                        f"'{label}' (row {row}) has no matching semester column."
                    )

                current_plan = {
                    "track": current_track,
                    "start": label,
                    "semesters": {},
                }
                current_start_column = start_column
                plans.append(current_plan)
            else:
                current_track = label
                current_plan = None
                current_start_column = None

        elif not cells:
            current_plan = None
            current_start_column = None

        for column, text in cells.items():
            if current_plan is None or column < current_start_column:
                unassigned_cells.append(
                    {"cell": sheet.cell(row, column).coordinate, "text": text}
                )
                continue

            number = column - current_start_column + 1
            semester = current_plan["semesters"].setdefault(number, {
                "number": number,
                "term_label": header_columns[column],
                "season": header_columns[column].rstrip("*").strip(),
                "entries": [],
            })
            semester["entries"].append(text)

    for plan in plans:
        plan["semesters"] = [plan["semesters"][n] for n in sorted(plan["semesters"])]

    if not plans:
        raise StudyPlanFormatError("No track / start-semester blocks were found in the study plan grid.")

    return plans, unassigned_cells, notes


def load_graduate_study_plan(path):
    workbook = _open_workbook(path)

    courses = parse_course_information(workbook)
    plans, unassigned_cells, notes = parse_study_plans(workbook)

    return {
        "courses": courses,
        "plans": plans,
        "unassigned_cells": unassigned_cells,
        "notes": notes,
    }


def get_plan(study_plan, track, start):
    wanted_track = _clean_text(track).lower()
    wanted_start = _clean_text(start).lower()

    for plan in study_plan["plans"]:
        if (_clean_text(plan["track"]).lower() == wanted_track
                and _clean_text(plan["start"]).lower() == wanted_start):
            return plan["semesters"]

    raise KeyError(f"No study plan found for track '{track}' with start '{start}'.")


def print_study_plan_results(study_plan):

    print("\nRECOMMENDED COURSE SEQUENCES")
    print("----------------------------")

    for plan in study_plan["plans"]:
        print(f"\n{plan['track']} | {plan['start']}")
        for semester in plan["semesters"]:
            print(
                f"  Semester {semester['number']} ({semester['term_label']}):",
                ", ".join(semester["entries"]),
            )

    print("\nCOURSE INFORMATION")
    print("------------------")

    for course in study_plan["courses"].values():
        print(
            course["course_code"],
            "|",
            course["course_name"],
            "| Prerequisite:",
            course["prerequisite"] or "None",
            "| Offered:",
            ", ".join(course["offered"]) or "Not listed",
        )

    if study_plan["unassigned_cells"]:
        print("\nCELLS OUTSIDE ANY TRACK BLOCK")
        print("-----------------------------")
        for item in study_plan["unassigned_cells"]:
            print(item["cell"], "|", item["text"])


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python src/GraduateStudyPlanParser.py <graduate_study_plans_xlsx>")
        sys.exit(1)

    study_plan_path = sys.argv[1]

    try:
        study_plan_data = load_graduate_study_plan(study_plan_path)
        print_study_plan_results(study_plan_data)

    except Exception as error:
        print(f"Error reading graduate study plans: {error}")
        sys.exit(1)
